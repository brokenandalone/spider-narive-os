"""Native writing tools ported from the Webbie Author Studio website.

The installed Author SQLite file remains authoritative. This widget neither
imports web-account data nor sends manuscripts to an external service.
"""
import difflib
import threading
from pathlib import Path
from PyQt5.QtCore import Qt, QThread, pyqtSignal
if __package__:
    from .publishing import export_pdf, export_epub
else:
    from publishing import export_pdf, export_epub
if __package__:
    from .voice_reader import WebbieReader
    from .review_engine import ReviewEngine, ReviewCancelled
else:
    from voice_reader import WebbieReader
    from review_engine import ReviewEngine, ReviewCancelled
from PyQt5.QtWidgets import (
    QComboBox, QFileDialog, QHBoxLayout, QLabel, QLineEdit, QListWidget,
    QListWidgetItem, QMessageBox, QPushButton, QTabWidget, QTextEdit,
    QVBoxLayout, QWidget
)


class ReviewWorker(QThread):
    progress = pyqtSignal(str)
    ready = pyqtSignal(str)
    failed = pyqtSignal(str)

    def __init__(self, chapters, title, scope, depth, canon, parent=None):
        super().__init__(parent)
        self.chapters = [dict(c) for c in chapters]
        self.title = title
        self.scope = scope
        self.depth = depth
        self.canon = canon
        self.cancelled = threading.Event()

    def cancel(self):
        self.cancelled.set()

    def run(self):
        try:
            report = ReviewEngine().review(
                self.chapters, self.title, scope=self.scope,
                depth=self.depth, canon=self.canon,
                progress=self.progress.emit, cancelled=self.cancelled,
            )
            if not self.cancelled.is_set():
                self.ready.emit(report)
            else:
                self.failed.emit('Review cancelled. Manuscript unchanged.')
        except ReviewCancelled:
            self.failed.emit('Review cancelled. Manuscript unchanged.')
        except Exception as error:
            self.failed.emit('Review could not finish: ' + str(error))


class AuthorToolkit(QWidget):
    jumpRequested = pyqtSignal(dict)
    reviewRequested = pyqtSignal(str)
    editRequested = pyqtSignal()

    def __init__(self, store, current_chapter, save_editor=None, editor=None):
        super().__init__()
        self.store = store
        self.current_chapter = current_chapter
        self.save_editor = save_editor
        self.editor = editor
        self.book_id = self.story_id = self.reference_id = None
        self.story_dirty = self.reference_dirty = False
        self.reader = WebbieReader(self)
        self.reader.changed.connect(self.read_status)
        self.review_worker = None
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel('BROKEN WORLD  /  Writing Desk'))
        tabs = QTabWidget(); layout.addWidget(tabs)
        search = QWidget(); s = QVBoxLayout(search)
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText('Find names, clues, chapters, canon and references across every book')
        self.search_input.returnPressed.connect(self.find)
        s.addWidget(self.search_input)
        find = QPushButton('Search Library'); find.clicked.connect(self.find); s.addWidget(find)
        self.search_results = QListWidget(); self.search_results.itemActivated.connect(self.jump)
        s.addWidget(self.search_results, 1)
        tabs.addTab(search, 'Search')
        story = QWidget(); b = QVBoxLayout(story)
        self.story_list = QListWidget(); b.addWidget(self.story_list, 1)
        self.story_category = QComboBox(); self.story_category.addItems(store.BIBLE_CATEGORIES)
        b.addWidget(self.story_category)
        self.story_title = QLineEdit(); self.story_title.setPlaceholderText('Character, location, timeline or canon rule')
        b.addWidget(self.story_title)
        self.story_body = QTextEdit(); self.story_body.setPlaceholderText('Established canon and continuity notes')
        b.addWidget(self.story_body, 2)
        story_row = QHBoxLayout(); b.addLayout(story_row)
        add = QPushButton('New Story Bible entry'); add.clicked.connect(self.new_story)
        story_row.addWidget(add)
        save = QPushButton('Save Story Bible'); save.clicked.connect(self.save_story)
        story_row.addWidget(save)
        self.story_list.currentItemChanged.connect(self.load_story)
        self.story_title.textChanged.connect(self.mark_story)
        self.story_category.currentIndexChanged.connect(self.mark_story)
        self.story_body.textChanged.connect(self.mark_story)
        tabs.addTab(story, 'Story Bible')
        references = QWidget(); r = QVBoxLayout(references)
        self.reference_list = QListWidget(); r.addWidget(self.reference_list, 1)
        self.reference_title = QLineEdit(); self.reference_title.setPlaceholderText('Companion reference title')
        r.addWidget(self.reference_title)
        self.reference_category = QLineEdit(); self.reference_category.setPlaceholderText('Category')
        r.addWidget(self.reference_category)
        self.reference_source = QLineEdit(); self.reference_source.setPlaceholderText('Original source name')
        r.addWidget(self.reference_source)
        self.reference_body = QTextEdit(); self.reference_body.setPlaceholderText('Source text, research and companion notes')
        r.addWidget(self.reference_body, 2)
        ref_row = QHBoxLayout(); r.addLayout(ref_row)
        add_ref = QPushButton('New reference'); add_ref.clicked.connect(self.new_reference)
        ref_row.addWidget(add_ref)
        save_ref = QPushButton('Save reference'); save_ref.clicked.connect(self.save_reference)
        ref_row.addWidget(save_ref)
        self.reference_list.currentItemChanged.connect(self.load_reference)
        for widget in (self.reference_title, self.reference_category, self.reference_source):
            widget.textChanged.connect(self.mark_reference)
        self.reference_body.textChanged.connect(self.mark_reference)
        tabs.addTab(references, 'Companion Library')
        tools = QWidget(); t = QVBoxLayout(tools)
        t.addWidget(QLabel('Exports and revision comparisons stay on your computer.'))
        t.addWidget(QLabel('On-demand Webbie reviews. No manuscript edits or cloud uploads.'))
        self.review_depth = QComboBox()
        self.review_depth.addItems(['Quick review', 'Deep review'])
        t.addWidget(self.review_depth)
        self.review_chapter_button = QPushButton('Webbie: Review entire chapter')
        self.review_chapter_button.clicked.connect(lambda: self.start_review('chapter'))
        t.addWidget(self.review_chapter_button)
        self.review_book_button = QPushButton('Webbie: Review entire book')
        self.review_book_button.clicked.connect(lambda: self.start_review('book'))
        t.addWidget(self.review_book_button)
        self.review_cancel_button = QPushButton('Cancel active review')
        self.review_cancel_button.clicked.connect(self.cancel_review)
        self.review_cancel_button.setEnabled(False)
        t.addWidget(self.review_cancel_button)
        self.review_status = QLabel('Review only starts when you request it.')
        self.review_status.setWordWrap(True)
        t.addWidget(self.review_status)
        self.review_result = QTextEdit()
        self.review_result.setReadOnly(True)
        self.review_result.setPlaceholderText('Full review report will appear here.')
        t.addWidget(self.review_result, 1)
        selected = QPushButton('Review selected passage with Webbie (preview first)')
        selected.clicked.connect(self.review_selection)
        t.addWidget(selected)
        approve = QPushButton('Apply reviewed rewrite (requires approval)')
        approve.clicked.connect(self.editRequested.emit)
        t.addWidget(approve)
        t.addWidget(QLabel('Narration uses Webbie’s configured voice, not the generic Author voice.'))
        self.read_button = QPushButton('Webbie: Read current chapter')
        self.read_button.clicked.connect(self.read_aloud)
        t.addWidget(self.read_button)
        self.read_book_button = QPushButton('Webbie: Read entire book')
        self.read_book_button.clicked.connect(self.read_book)
        t.addWidget(self.read_book_button)
        self.pause_button = QPushButton('Pause narration after current section')
        self.pause_button.clicked.connect(self.reader.pause)
        t.addWidget(self.pause_button)
        self.resume_button = QPushButton('Resume narration')
        self.resume_button.clicked.connect(self.reader.resume)
        t.addWidget(self.resume_button)
        self.stop_button = QPushButton('Stop narration')
        self.stop_button.clicked.connect(self.reader.stop)
        t.addWidget(self.stop_button)
        self.read_label = QLabel('Webbie voice: Natasha when available; existing Webbie fallback offline.')
        self.read_label.setWordWrap(True)
        t.addWidget(self.read_label)
        self.stats = QLabel('No project selected'); t.addWidget(self.stats)
        export = QPushButton('Export current book to DOCX'); export.clicked.connect(self.export_docx)
        t.addWidget(export)
        for label, extension, handler in (
            ('Export ebook (EPUB 3)', 'epub', export_epub),
            ('Export reading PDF', 'pdf', export_pdf),
        ):
            btn = QPushButton(label)
            btn.clicked.connect(lambda checked=False, suffix=extension, task=handler: self.export_other(suffix, task))
            t.addWidget(btn)
        compare = QPushButton('Compare current chapter with last saved version'); compare.clicked.connect(self.compare)
        t.addWidget(compare)
        self.compare_result = QTextEdit(); self.compare_result.setReadOnly(True)
        self.compare_result.setPlaceholderText('Revisions are shown here without changing the manuscript.')
        t.addWidget(self.compare_result, 1)
        tabs.addTab(tools, 'Publishing & Revisions')
        self.set_book(None)

    def mark_story(self, *_):
        if self.story_id is not None: self.story_dirty = True

    def mark_reference(self, *_):
        if self.reference_id is not None: self.reference_dirty = True

    def save_current(self):
        # Called by the parent editor autosave and before switching chapters.
        if self.story_dirty and self.story_id is not None:
            self.store.save_story_entry(self.story_id, self.story_category.currentText(),
                                        self.story_title.text(), self.story_body.toPlainText())
            self.story_dirty = False
        if self.reference_dirty and self.reference_id is not None:
            self.store.save_reference(self.reference_id, self.reference_title.text(),
                                      self.reference_category.text(), self.reference_body.toPlainText(),
                                      self.reference_source.text())
            self.reference_dirty = False
        return True

    def set_book(self, book_id):
        self.save_current()
        self.book_id = book_id
        self.story_id = self.reference_id = None
        self.refresh_story()
        self.refresh_references()
        self.refresh_stats()

    def refresh_stats(self):
        if self.book_id is not None:
            self.stats.setText(str(self.store.word_count(self.book_id)) + ' words in this book')
        else:
            self.stats.setText('Choose a book first.')

    def refresh_story(self, select=None):
        self.story_list.blockSignals(True); self.story_list.clear()
        if self.book_id is not None:
            for entry in self.store.story_entries(self.book_id):
                item = QListWidgetItem(entry['category'] + ' / ' + entry['title'])
                item.setData(Qt.UserRole, entry['id']); self.story_list.addItem(item)
                if entry['id'] == select: self.story_list.setCurrentItem(item)
        self.story_list.blockSignals(False)
        self.load_story(self.story_list.currentItem())

    def load_story(self, item, previous=None):
        try: self.save_current()
        except Exception as exc:
            QMessageBox.warning(self, 'Story Bible', str(exc))
            self.story_list.blockSignals(True); self.story_list.setCurrentItem(previous); self.story_list.blockSignals(False)
            return
        self.story_id = item.data(Qt.UserRole) if item else None
        entry = next((r for r in self.store.story_entries(self.book_id) if r['id']==self.story_id), None) if self.book_id is not None else None
        self.story_title.setText(entry['title'] if entry else '')
        self.story_body.setPlainText(entry['body'] if entry else '')
        self.story_category.setCurrentText(entry['category'] if entry else 'Canon')
        self.story_dirty = False
        for widget in (self.story_title, self.story_body, self.story_category):
            widget.setEnabled(entry is not None)

    def new_story(self):
        if self.book_id is None: return
        try:
            self.save_current()
            ident = self.store.add_story_entry(self.book_id, 'Canon', 'New canon entry')
            self.refresh_story(select=ident)
        except Exception as exc: QMessageBox.warning(self, 'Story Bible', str(exc))

    def save_story(self):
        try: self.save_current()
        except Exception as exc: QMessageBox.warning(self, 'Story Bible', str(exc))

    def refresh_references(self, select=None):
        self.reference_list.blockSignals(True); self.reference_list.clear()
        if self.book_id is not None:
            for entry in self.store.references(self.book_id):
                item = QListWidgetItem(entry['title']); item.setData(Qt.UserRole, entry['id'])
                self.reference_list.addItem(item)
                if entry['id'] == select: self.reference_list.setCurrentItem(item)
        self.reference_list.blockSignals(False)
        self.load_reference(self.reference_list.currentItem())

    def load_reference(self, item, previous=None):
        try: self.save_current()
        except Exception as exc:
            QMessageBox.warning(self, 'Companion Library', str(exc))
            self.reference_list.blockSignals(True); self.reference_list.setCurrentItem(previous); self.reference_list.blockSignals(False)
            return
        self.reference_id = item.data(Qt.UserRole) if item else None
        entry = next((r for r in self.store.references(self.book_id) if r['id']==self.reference_id), None) if self.book_id is not None else None
        self.reference_title.setText(entry['title'] if entry else '')
        self.reference_category.setText(entry['category'] if entry else '')
        self.reference_source.setText(entry['source_name'] if entry else '')
        self.reference_body.setPlainText(entry['body'] if entry else '')
        self.reference_dirty = False
        for widget in (self.reference_title, self.reference_category, self.reference_source, self.reference_body):
            widget.setEnabled(entry is not None)

    def new_reference(self):
        if self.book_id is None: return
        try:
            self.save_current()
            ident = self.store.add_reference(self.book_id, 'New companion reference')
            self.refresh_references(select=ident)
        except Exception as exc: QMessageBox.warning(self, 'Companion Library', str(exc))

    def save_reference(self):
        try: self.save_current()
        except Exception as exc: QMessageBox.warning(self, 'Companion Library', str(exc))

    def find(self):
        self.search_results.clear()
        try: hits = self.store.search_library(self.search_input.text())
        except Exception as exc:
            QMessageBox.warning(self, 'Search', str(exc)); return
        for hit in hits:
            item = QListWidgetItem(hit['book_title'] + ' / ' + hit['kind'] + ' / ' + hit['title'])
            item.setToolTip(hit['excerpt']); item.setData(Qt.UserRole, hit)
            self.search_results.addItem(item)

    def jump(self, item):
        try: self.save_current()
        except Exception as exc:
            QMessageBox.warning(self, 'Search', str(exc)); return
        self.jumpRequested.emit(item.data(Qt.UserRole))

    def compare(self):
        self.compare_result.clear()
        if self.save_editor is not None and not self.save_editor(): return
        chapter = self.current_chapter()
        if chapter is None: return
        snapshots = self.store.snapshots(chapter['id'])
        if not snapshots:
            self.compare_result.setPlainText('No previous chapter version yet.'); return
        old = snapshots[0]['content'].splitlines(keepends=True)
        new = chapter['content'].splitlines(keepends=True)
        diff = difflib.unified_diff(old, new, fromfile='Earlier version', tofile='Current version')
        self.compare_result.setPlainText(''.join(diff) or 'No text changes.')

    def read_status(self, message):
        self.read_label.setText(message)

    def read_aloud(self):
        if self.save_editor is not None and not self.save_editor(): return
        chapter = self.current_chapter()
        if chapter is None:
            QMessageBox.information(self, 'Webbie voice', 'Select a chapter to read.')
            return
        try:
            self.reader.start_chapter(chapter['content'], chapter['title'])
        except (OSError, RuntimeError, ValueError) as exc:
            QMessageBox.warning(self, 'Webbie voice', str(exc))

    def read_book(self):
        if self.book_id is None:
            QMessageBox.information(self, 'Webbie voice', 'Select a book to read.')
            return
        if self.save_editor is not None and not self.save_editor(): return
        try:
            chapters = [dict(c) for c in self.store.chapters(self.book_id)]
            self.reader.start_book(chapters)
        except (OSError, RuntimeError, ValueError) as exc:
            QMessageBox.warning(self, 'Webbie voice', str(exc))

    def start_review(self, scope):
        if self.review_worker is not None and self.review_worker.isRunning():
            return
        if self.book_id is None:
            QMessageBox.information(self, 'Webbie review', 'Select a book first.')
            return
        if scope == 'chapter' and self.current_chapter() is None:
            QMessageBox.information(self, 'Webbie review', 'Select a chapter first.')
            return
        if self.save_editor is not None and not self.save_editor():
            return
        book = next((b for b in self.store.books() if b['id'] == self.book_id), None)
        if book is None:
            return
        chapters = ([dict(self.current_chapter())] if scope == 'chapter'
                    else [dict(c) for c in self.store.chapters(self.book_id)])
        if not any(str(c['content']).strip() for c in chapters):
            QMessageBox.information(self, 'Webbie review', 'No manuscript text to review.')
            return
        depth = 'deep' if self.review_depth.currentIndex() else 'quick'
        self.review_worker = ReviewWorker(chapters, str(book['title']), scope, depth,
                                          self.store.canon(self.book_id), self)
        self.review_worker.progress.connect(self.review_status.setText)
        self.review_worker.ready.connect(self.review_finished_report)
        self.review_worker.failed.connect(self.review_status.setText)
        self.review_worker.finished.connect(self.review_finished)
        self.review_result.setPlainText('Review running. Every section will be inspected. '
                                        'Manuscript remains unchanged.')
        self.review_chapter_button.setEnabled(False)
        self.review_book_button.setEnabled(False)
        self.review_cancel_button.setEnabled(True)
        self.review_worker.start()

    def cancel_review(self):
        if self.review_worker and self.review_worker.isRunning():
            self.review_worker.cancel()
            self.review_status.setText('Cancelling after the current local model request…')

    def review_finished(self):
        self.review_chapter_button.setEnabled(True)
        self.review_book_button.setEnabled(True)
        self.review_cancel_button.setEnabled(False)

    def review_finished_report(self, report):
        self.review_result.setPlainText(report)
        self.review_status.setText('Review complete. No changes were made to your book.')

    def shutdown(self):
        # Keep QObject/QThread alive until a pending inference finishes.
        if self.review_worker and self.review_worker.isRunning():
            self.cancel_review()
            return False
        self.reader.stop()
        return self.reader.wait(3500) if self.reader.isRunning() else True

    def review_selection(self):
        # The user must explicitly click this button and then press Send
        # in Webbie's sidebar. Never automatically transmit chapter text.
        editor = self.editor
        chapter = self.current_chapter()
        title = chapter['title'] if chapter is not None else 'Untitled'
        selection = editor.textCursor().selectedText().replace('\u2029', '\n') if editor else ''
        selection = selection.strip()
        if not selection:
            QMessageBox.information(self, 'Webbie review', 'Highlight a passage in the chapter first.')
            return
        if len(selection) > 4000:
            QMessageBox.warning(self, 'Webbie review', 'Select at most 4,000 characters at a time.')
            return
        self.reviewRequested.emit(
            'Writer requests a continuity and prose review for "' + title +
            '". Suggest improvements but do not edit the manuscript.\n\n' + selection
        )

    def export_other(self, suffix, task):
        if self.book_id is None:
            return
        if self.save_editor is not None and not self.save_editor():
            return
        try:
            self.save_current()
            book = next(b for b in self.store.books() if b['id'] == self.book_id)
            destination, _ = QFileDialog.getSaveFileName(
                self, 'Export manuscript as ' + suffix.upper(),
                book['title'] + '.' + suffix,
                suffix.upper() + ' (*.' + suffix + ')'
            )
            if destination:
                task(self.store, self.book_id, destination)
        except Exception as exc:
            QMessageBox.warning(self, 'Export', str(exc))

    def export_docx(self):
        if self.book_id is None: return
        if self.save_editor is not None and not self.save_editor(): return
        try:
            self.save_current()
            book = next(b for b in self.store.books() if b['id'] == self.book_id)
            destination, _ = QFileDialog.getSaveFileName(
                self, 'Export manuscript', book['title'] + '.docx', 'Word document (*.docx)'
            )
            if not destination: return
            from docx import Document
            document = Document()
            document.add_heading(book['title'], 0)
            for chapter in self.store.chapters(self.book_id):
                document.add_page_break()
                document.add_heading(chapter['title'], level=1)
                for paragraph in chapter['content'].split('\n'):
                    document.add_paragraph(paragraph)
            with open(destination, 'xb') as output:
                document.save(output)
        except Exception as exc: QMessageBox.warning(self, 'Export', str(exc))
