"""Build user-requested APA student papers; no content or citation invention."""
import os
from pathlib import Path
import tempfile


def create_paper(path, *, title, author, institution, course, instructor, due,
                 paragraphs=(), references=(), overwrite=False):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    if any(not str(value).strip() for value in (title, author, institution, course, instructor, due)):
        raise ValueError('Complete all title-page fields.')
    path = Path(path)
    if path.suffix.lower() != '.docx':
        raise ValueError('Choose a DOCX filename.')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(.5)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Inches(0)
    run = header.add_run(); field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run._r.addnext(field)
    def centered(text, bold=False):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Inches(0)
        p.add_run(text).bold=bold
        return p
    for _ in range(3): centered('')
    centered(title, True); centered('')
    for value in (author, institution, course, instructor, due): centered(value)
    doc.add_page_break()
    centered(title, True)
    for text in paragraphs:
        doc.add_paragraph(text)
    if not paragraphs: doc.add_paragraph('')
    doc.add_page_break(); centered('References', True)
    for text in references:
        p=doc.add_paragraph(text)
        p.paragraph_format.left_indent=Inches(.5)
        p.paragraph_format.first_line_indent=Inches(-.5)
    if not references:
        p=doc.add_paragraph('')
        p.paragraph_format.left_indent=Inches(.5)
        p.paragraph_format.first_line_indent=Inches(-.5)
    # Build in the same filesystem, then publish only a complete document.
    handle, temporary = tempfile.mkstemp(prefix='.apa-', suffix='.docx', dir=path.parent)
    os.close(handle)
    try:
        doc.save(temporary)
        with open(temporary, 'rb') as stream: os.fsync(stream.fileno())
        if overwrite:
            os.replace(temporary, path)
        else:
            os.link(temporary, path)  # Refuses to clobber an existing file, including races.
    finally:
        Path(temporary).unlink(missing_ok=True)
    return path
