"""Read text from a school document the owner explicitly selects.

Never crawl OneDrive or silently add files to a prompt. PDF text extraction
uses poppler's pdftotext; image-only PDFs require a separate owner decision.
"""
import shutil
import subprocess
from pathlib import Path

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_TEXT_CHARS = 12000
SUPPORTED = frozenset({".txt", ".md", ".docx", ".pdf"})


def safe_remote_file(folder, name):
    """Validate a single chosen entry from the OneDrive folder listing."""
    parts = str(folder).split("/") if folder else []
    parts.append(str(name))
    if any(not part or part in (".", "..") or
           "\\" in part or "\x00" in part or
           any(ord(char) < 32 for char in part) for part in parts):
        raise ValueError("Unsafe OneDrive file path.")
    if "/" in str(name) or ":" in parts[0] or str(name).startswith("-"):
        raise ValueError("Unsafe OneDrive file name.")
    if Path(str(name)).suffix.lower() not in SUPPORTED:
        raise ValueError("Select a TXT, Markdown, Word DOCX, or text PDF.")
    return "/".join(parts)


def read_course_document(path):
    """Extract bounded text from an explicit local or downloaded document."""
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise ValueError("Choose an existing regular document, not a link.")
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError("Supported school files: TXT, MD, DOCX, and text PDF.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("This document exceeds the 8 MB import limit.")
    if path.suffix.lower() in (".txt", ".md"):
        text = path.read_text(encoding="utf-8-sig")
    elif path.suffix.lower() == ".docx":
        from docx import Document
        document = Document(path)
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(paragraphs)
    else:
        if not shutil.which("pdftotext"):
            raise RuntimeError("PDF reading requires poppler-utils (pdftotext).")
        result = subprocess.run(
            ["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
            capture_output=True, text=True, timeout=25, check=False)
        if result.returncode != 0:
            raise ValueError("Unable to read this PDF's text.")
        text = result.stdout
    text = text.strip()
    if not text:
        raise ValueError("No readable text found. Scanned or image-only files need OCR.")
    return text[:MAX_TEXT_CHARS]


def format_selected_source(remote_file, text):
    """Preserve origin; never turn document text into verified citations."""
    if not isinstance(text, str) or not text.strip():
        raise ValueError("The selected document contains no usable text.")
    return ("SCHOOL ONEDRIVE FILE: " + remote_file +
            "\nUnverified user-selected course document (not a citation):\n" +
            text[:MAX_TEXT_CHARS])
