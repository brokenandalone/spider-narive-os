import importlib.util
from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
def load(name):
    spec=importlib.util.spec_from_file_location('school_'+name,ROOT/'study'/f'{name}.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module);return module
apa=load('apa');dashboard=load('dashboard');store_module=load('store')

FIELDS=dict(title='Paper Title',author='Student Name',institution='University',course='Course',instructor='Instructor',due='October 8, 2026')

class SchoolToolsTests(unittest.TestCase):
    def test_dates_keep_unparseable_values_visible_and_completed_out_of_overdue(self):
        rows=[{'title':title,'due':due,'status':status,'course':'Course'} for title,due,status in [
            ('Old','2026-10-01','Open'),('Done','2026-09-01','Complete'),
            ('Next','2026-10-09','Open'),('Unknown','Sunday','Open')]]
        result=dashboard.overview(rows,date(2026,10,8))
        self.assertEqual((result['open'],result['complete'],result['overdue'],result['undated']),(3,1,1,1))
        self.assertEqual([r['title'] for r in result['upcoming']],['Old','Next'])

    def test_paper_preserves_supplied_content_and_formatting(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'paper.docx'
            apa.create_paper(path,**FIELDS,paragraphs=['First paragraph.','Second paragraph.'],references=['Author. (2026). Source.'])
            doc=Document(path);section=doc.sections[0]
            self.assertEqual(section.left_margin,Inches(1))
            normal=doc.styles['Normal']
            self.assertEqual(normal.font.name,'Times New Roman');self.assertEqual(normal.font.size,Pt(12))
            self.assertEqual(normal.paragraph_format.line_spacing,2)
            self.assertEqual(normal.paragraph_format.first_line_indent,Inches(.5))
            self.assertIn('First paragraph.',[p.text for p in doc.paragraphs])
            reference=next(p for p in doc.paragraphs if p.text=='Author. (2026). Source.')
            self.assertEqual(reference.paragraph_format.left_indent,Inches(.5))
            self.assertEqual(reference.paragraph_format.first_line_indent,Inches(-.5))
            self.assertIn('PAGE',section.header._element.xml)

    def test_existing_document_and_failed_export_are_preserved(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'paper.docx';path.write_bytes(b'Original')
            with self.assertRaises(FileExistsError):apa.create_paper(path,**FIELDS)
            self.assertEqual(path.read_bytes(),b'Original')
            with patch('docx.document.Document.save',side_effect=OSError('disk full')):
                with self.assertRaises(OSError):apa.create_paper(path,**FIELDS,overwrite=True)
            self.assertEqual(path.read_bytes(),b'Original')
            self.assertEqual(list(Path(folder).glob('.apa-*')),[])

    def test_dashboard_reads_existing_database_without_schema_or_path_migration(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(store_module,'DATA_DIR',Path(folder)/'data'), \
                 patch.object(store_module,'DB_PATH',Path(folder)/'data/study.db'), \
                 patch.object(store_module,'DOCUMENTS_ROOT',Path(folder)/'Study'):
                store=store_module.StudyStore();course=store.add_course('Course')
                store.add_assignment(course,'Essay','2026-10-09');store.save_notes(course,'Keep notes')
                store.close();store=store_module.StudyStore()
                self.assertEqual(store.all_assignments()[0]['course'],'Course')
                self.assertEqual(store.get_notes(course),'Keep notes');store.close()
